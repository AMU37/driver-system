"""تحديث «خط السير» للموظفين من كشف «بأسماء ومواقع السكن وخطوط سير الموظفين».

يحدّث عمود route_name فقط حسب كود الموظف، ولا يمسّ أي حقل آخر.
التشغيل (من مجلد backend) على القاعدة المحلية:
    .venv\\Scripts\\python.exe scripts\\import_routes.py [--path <file.xlsx>] [--dry-run]
والتشغيل على قاعدة الإنتاج (Supabase/Postgres) بتمرير رابط القاعدة:
    .venv\\Scripts\\python.exe scripts\\import_routes.py --database-url "postgresql://...?sslmode=require"
"""
import argparse
import os
import sys

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.core.database import SessionLocal  # noqa: E402
from app.models import Employee  # noqa: E402

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_FILE = os.path.join(
    BACKEND_DIR, "..", "كشف بأسماء ومواقع السكن وخطوط سير الموظفين.xlsx"
)
CODE_HEADERS = {"كود الموظف", "كودالموظف", "الكود"}
ROUTE_HEADERS = {"خط السير", "خطالسير", "خط سير"}


def _clean_code(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def read_routes(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header = None
    for row in rows:
        cells = [("" if c is None else str(c)).strip() for c in (row or [])]
        if any(c in CODE_HEADERS for c in cells) and any(c in ROUTE_HEADERS for c in cells):
            header = cells
            break
    if header is None:
        raise SystemExit("تعذر العثور على صف الترويسة (كود الموظف / خط السير)")

    ci = next(i for i, c in enumerate(header) if c in CODE_HEADERS)
    ri = next(i for i, c in enumerate(header) if c in ROUTE_HEADERS)

    start = rows.index(next(r for r in rows if [("" if c is None else str(c)).strip() for c in (r or [])] == header)) + 1
    mapping = {}
    for row in rows[start:]:
        if not row:
            continue
        code = _clean_code(row[ci] if ci < len(row) else None)
        if not code:
            continue
        route = ""
        if ri < len(row) and row[ri] is not None:
            route = str(row[ri]).strip()
        mapping[code] = route
    return mapping


def _make_session(database_url):
    url = database_url.strip()
    if url.startswith("postgres://"):
        url = "postgresql+psycopg://" + url[len("postgres://"):]
    elif url.startswith("postgresql://"):
        url = "postgresql+psycopg://" + url[len("postgresql://"):]
    engine = create_engine(url, pool_pre_ping=True)
    if engine.dialect.name == "postgresql":
        with engine.begin() as conn:
            conn.exec_driver_sql("ALTER TABLE employees ADD COLUMN IF NOT EXISTS route_name VARCHAR(200)")
    return sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--path", default=DEFAULT_FILE)
    parser.add_argument("--database-url", default=None, help="رابط قاعدة البيانات (اختياري، للإنتاج)")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    path = os.path.abspath(args.path)
    if not os.path.exists(path):
        raise SystemExit(f"الملف غير موجود: {path}")

    mapping = read_routes(path)
    with_value = {c: r for c, r in mapping.items() if r and r != "#N/A"}

    Session = _make_session(args.database_url) if args.database_url else SessionLocal
    db = Session()
    try:
        employees = db.scalars(select(Employee)).all()
        by_code = {e.employee_code: e for e in employees}

        updated = skipped_empty = unmatched = already = 0
        for code, route in with_value.items():
            employee = by_code.get(code)
            if not employee:
                unmatched += 1
                continue
            if (employee.route_name or "") == route:
                already += 1
                continue
            employee.route_name = route
            updated += 1

        skipped_empty = len(mapping) - len(with_value)
        if not args.dry_run:
            db.commit()

        print(f"file={path}")
        print(f"rows={len(mapping)} with_route={len(with_value)} empty_or_na={skipped_empty}")
        print(f"updated={updated} already_same={already} unmatched_code={unmatched} total_employees={len(employees)}")
        if args.dry_run:
            print("(dry-run — لم تُحفظ التغييرات)")
    finally:
        db.close()


if __name__ == "__main__":
    main()
