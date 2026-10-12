import csv
import io
from datetime import date, datetime, time

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Employee

_ALIASES = {
    "employeecode": "employee_code",
    "code": "employee_code",
    "كودالموظف": "employee_code",
    "الكود": "employee_code",
    "name": "name",
    "fullname": "name",
    "الاسم": "name",
    "اسمالموظف": "name",
    "الاسمالرسمي": "name",
    "jobtitle": "job_title",
    "job": "job_title",
    "الوظيفة": "job_title",
    "المسمي": "job_title",
    "المسمى": "job_title",
    "المسمىالوظيفي": "job_title",
    "departmentname": "department_name",
    "department": "department_name",
    "الادارة": "department_name",
    "الإدارة": "department_name",
    "القسم": "department_name",
    "اسمادارةالفرعية": "department_name",
    "اسمالإدارةالفرعي": "department_name",
    "إسمالادارةالفرعية": "department_name",
    "إسمالإدارةالفرعية": "department_name",
    "companyname": "company_name",
    "company": "company_name",
    "الشركة": "company_name",
    "housinglocation": "housing_location",
    "housing": "housing_location",
    "السكن": "housing_location",
    "موقعالسكن": "housing_location",
    # خطة الدوام والدورة الفعلية من كشف «الموظفين مع نظام الدوام»
    "جدولالعمل": "work_schedule",
    "شرحنوعالدوام": "schedule_desc",
    "شرحأنواعالدوام": "schedule_desc",
    "عددأيامالعمل": "work_days",
    "عددأيامالراحة": "rest_days",
    "تاريخاوليومعمل": "first_work_day",
    "تاريخاوليومعمل": "first_work_day",
    "اوليومعمل": "first_work_day",
    "أوليومعمل": "first_work_day",
    "الحالةالدوريةللموظف": "rotation_status",
    "الحالةالدورية": "rotation_status",
    "وقتالمغادرة": "departure_time",
    "نوعالجنس": "gender",
    "الجنس": "gender",
    "نوعالخدمة": "service_type",
    "تصنيفالنقل": "transport_type",
    "نقلدوري": "transport_type",
    "دوري": "transport_type",
    "يومالدورةاليوم": "cycle_day_today",
    "دورةاليوم": "cycle_day_today",
    "الحالةالأساسيةاليوم": "base_status_today",
    "الحالةالاساسيةاليوم": "base_status_today",
    "يومالدورةغدا": "cycle_day_tomorrow",
    "يومالدورةغداً": "cycle_day_tomorrow",
    "دورةغدا": "cycle_day_tomorrow",
    "دورةغداً": "cycle_day_tomorrow",
    # خط السير من كشف «بأسماء ومواقع السكن وخطوط سير الموظفين»
    "خطالسير": "route_name",
    "خطسير": "route_name",
    "routename": "route_name",
    "route": "route_name",
}

_FIELDS = [
    "employee_code", "name", "job_title", "department_name", "company_name", "housing_location",
    "work_schedule", "schedule_desc", "work_days", "rest_days", "first_work_day",
    "rotation_status", "departure_time", "gender", "service_type", "transport_type",
    "cycle_day_today", "base_status_today", "cycle_day_tomorrow", "route_name",
]

MAX_IMPORT_ROWS = 50000
MAX_FIELD_LEN = 200

_NUMBER_FIELDS = {"work_days", "rest_days", "cycle_day_today", "cycle_day_tomorrow"}


def _to_int(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value).strip()
    if not text:
        return None
    digits = "".join(ch for ch in text if ch.isdigit())
    return int(digits) if digits else None


def _sanitize(text: str) -> str:
    # Neutralize CSV/Excel formula injection: values beginning with = + - @
    # or tab/CR are prefixed with a single quote so they are never evaluated
    # as formulas when the data is later opened in a spreadsheet.
    if text and (text[0] in "=+-@" or text[:1] in ("\t", "\r")):
        return "'" + text
    if len(text) > MAX_FIELD_LEN:
        return text[:MAX_FIELD_LEN]
    return text


def _norm_header(value: str) -> str:
    return "".join(ch for ch in value.lower() if ch.isalnum())


def _map_columns(headers: list[str]) -> list[tuple[str, str]]:
    used: set[str] = set()
    mapped: list[tuple[str, str]] = []
    for header in headers:
        field = _ALIASES.get(_norm_header(header or ""))
        if field and field not in used:
            mapped.append((header, field))
            used.add(field)
    return mapped


def _clean(value):
    if value is None:
        return None
    # خلايا Excel للتاريخ/الوقت تُحوَّل إلى صيغة نصية قياسية
    # (YYYY-MM-DD و HH:MM:SS) ليقرأها التطبيق ويحسب الدورة بثبات.
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return _sanitize(str(int(value)))
    text = str(value).strip()
    return _sanitize(text) if text else None


def parse_employee_rows(data: bytes, filename: str) -> list[dict]:
    filename = (filename or "").lower()
    if filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = None
        mapping: list[tuple[str, str]] = []
        for raw in rows:
            candidate = [str(c) if c is not None else "" for c in raw]
            if not any(cell.strip() for cell in candidate):
                continue
            found = _map_columns(candidate)
            if found:
                header = candidate
                mapping = found
                break
        if header is None:
            workbook.close()
            return []
        parsed = []
        for record in rows:
            if record is None or all(cell is None or str(cell).strip() == "" for cell in record):
                continue
            if len(parsed) >= MAX_IMPORT_ROWS:
                raise ValueError(f"الملف يحتوي أكثر من {MAX_IMPORT_ROWS} صف")
            row: dict[str, str] = {}
            for idx, field in mapping:
                col = header.index(idx)
                if col < len(record):
                    row[field] = _clean(record[col])
            parsed.append(row)
        workbook.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("الملف فارغ أو بدون ترويسة")
        mapping = _map_columns(reader.fieldnames)
        parsed = []
        for record in reader:
            if any(str(v or "").strip() for v in record.values()):
                if len(parsed) >= MAX_IMPORT_ROWS:
                    raise ValueError(f"الملف يحتوي أكثر من {MAX_IMPORT_ROWS} صف")
                row: dict[str, str] = {}
                for header, field in mapping:
                    row[field] = _clean(record.get(header))
                parsed.append(row)
    return [row for row in parsed if row.get("employee_code") or row.get("name")]


def import_employee_rows(db: Session, rows: list[dict]) -> dict:
    created = updated = 0
    errors: list[dict] = []
    for idx, row in enumerate(rows, start=1):
        employee_code = _clean(row.get("employee_code"))
        name = _clean(row.get("name"))
        if not employee_code and not name:
            continue
        if not employee_code or not name:
            errors.append({"row": idx, "error": "يجب إدخال كود الموظف واسمه"})
            continue
        employee = db.scalar(select(Employee).where(Employee.employee_code == employee_code))
        if not employee:
            employee = Employee(employee_code=employee_code, name=name)
            db.add(employee)
            created += 1
        else:
            updated += 1
        employee.name = name
        employee.job_title = _clean(row.get("job_title")) or employee.job_title
        employee.department_name = _clean(row.get("department_name")) or employee.department_name
        employee.company_name = _clean(row.get("company_name")) or employee.company_name
        employee.housing_location = _clean(row.get("housing_location")) or employee.housing_location
        employee.work_schedule = _clean(row.get("work_schedule")) or employee.work_schedule
        employee.schedule_desc = _clean(row.get("schedule_desc")) or employee.schedule_desc
        employee.first_work_day = _clean(row.get("first_work_day")) or employee.first_work_day
        employee.rotation_status = _clean(row.get("rotation_status")) or employee.rotation_status
        employee.departure_time = _clean(row.get("departure_time")) or employee.departure_time
        employee.gender = _clean(row.get("gender")) or employee.gender
        employee.service_type = _clean(row.get("service_type")) or employee.service_type
        employee.transport_type = _clean(row.get("transport_type")) or employee.transport_type
        employee.base_status_today = _clean(row.get("base_status_today")) or employee.base_status_today
        employee.route_name = _clean(row.get("route_name")) or employee.route_name
        for field in _NUMBER_FIELDS:
            value = _to_int(row.get(field))
            setattr(employee, field, value if value is not None else getattr(employee, field))
        employee.is_active = True
    db.commit()
    return {"created": created, "updated": updated, "skipped": errors}